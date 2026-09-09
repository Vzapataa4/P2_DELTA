"""Importaciones asistidas por archivo (mejoras pedidas por Mauricio y
detectadas en la revisión de Sprint 0):

1. Foto de factura -> pre-llenar el registro de venta (Acta de Compromiso,
   5 de agosto de 2026: "en ocasiones se agotan ciertos insumos... la
   información se encuentra muy dispersa").
2. Excel de receta -> pre-llenar los insumos de una receta.

Ninguna de las dos funciones guarda datos directamente: ambas devuelven lo
que lograron leer para que el usuario confirme o corrija antes de guardar.
Esto es intencional -- en Sprint 0 encontramos un Excel real cuyo nombre
interno no coincidía con el producto real, así que nunca se debe confiar
ciegamente en un archivo importado (ver Wiki_RNF_Ampliados.md, Manejo de
fallos).
"""
import base64
import json
import logging
import os
import re
from io import BytesIO
from typing import List

from openpyxl import load_workbook

from . import models
from .schemas import (
    ExcelInsumoLineaOut,
    ExcelRecetaImportOut,
    FacturaImportOut,
    FacturaLineaOut,
)

logger = logging.getLogger(__name__)

ANTHROPIC_API_KEY = os.getenv("ANTHROPIC_API_KEY")
CLAUDE_MODEL = os.getenv("CLAUDE_MODEL", "claude-sonnet-5")
IA_TIMEOUT_SECONDS = 15

_MEDIA_TYPES_VALIDOS = {"image/jpeg", "image/png", "image/webp", "image/gif"}


def extraer_factura_desde_imagen(
    contenido: bytes, content_type: str, recetas_conocidas: List[str]
) -> FacturaImportOut:
    """Lee una foto de factura/pedido y devuelve distribuidor + líneas de
    producto sugeridas. Si no hay ANTHROPIC_API_KEY configurada, o la
    lectura falla, se degrada honestamente: no se inventan productos, se
    le pide al usuario completar el formulario a mano."""
    if not ANTHROPIC_API_KEY:
        logger.info("ANTHROPIC_API_KEY no configurada; no se puede leer la factura automáticamente.")
        return FacturaImportOut(
            distribuidor_nombre="",
            lineas=[],
            fuente="manual",
            mensaje=(
                "La lectura automática de facturas no está activada en este ambiente "
                "(falta ANTHROPIC_API_KEY). Completa el distribuidor y los productos manualmente."
            ),
        )

    try:
        import anthropic  # import perezoso: solo si la key está configurada

        client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY, timeout=IA_TIMEOUT_SECONDS)
        imagen_b64 = base64.b64encode(contenido).decode("utf-8")
        media_type = content_type if content_type in _MEDIA_TYPES_VALIDOS else "image/jpeg"

        prompt = (
            "Esta imagen es la foto de una factura o pedido de una panadería. "
            "Extrae el nombre del distribuidor/cliente y cada producto con su cantidad en unidades. "
            "Usa exactamente uno de estos nombres de producto cuando alguno coincida o se parezca: "
            f"{', '.join(recetas_conocidas)}. "
            "Responde ÚNICAMENTE con JSON válido, sin texto adicional ni explicación, con esta forma exacta: "
            '{"distribuidor_nombre": "...", "lineas": [{"receta_nombre": "...", "cantidad": 0}]}'
        )

        resp = client.messages.create(
            model=CLAUDE_MODEL,
            max_tokens=500,
            messages=[
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "image",
                            "source": {"type": "base64", "media_type": media_type, "data": imagen_b64},
                        },
                        {"type": "text", "text": prompt},
                    ],
                }
            ],
        )
        texto = resp.content[0].text.strip()
        texto = re.sub(r"^```(json)?|```$", "", texto, flags=re.MULTILINE).strip()
        datos = json.loads(texto)

        lineas = [
            FacturaLineaOut(receta_nombre=str(linea["receta_nombre"]), cantidad=int(linea["cantidad"]))
            for linea in datos.get("lineas", [])
            if linea.get("cantidad") and int(linea["cantidad"]) > 0
        ]
        return FacturaImportOut(
            distribuidor_nombre=str(datos.get("distribuidor_nombre") or ""),
            lineas=lineas,
            fuente="ia",
            mensaje="Factura leída automáticamente. Revisa los datos antes de confirmar la venta.",
        )
    except Exception as exc:  # nunca tumbar la funcionalidad por un fallo de lectura
        logger.warning("Fallo al leer la factura con IA, se pide llenado manual: %s", exc)
        return FacturaImportOut(
            distribuidor_nombre="",
            lineas=[],
            fuente="manual",
            mensaje="No se pudo leer la factura automáticamente. Completa los datos manualmente.",
        )


def leer_excel_receta(contenido: bytes, nombre_archivo: str, db) -> ExcelRecetaImportOut:
    """Lee un Excel con insumos y cantidades de una receta. No guarda nada
    directamente -- el bug real de Sprint 0 fue justamente un Excel cuya
    etiqueta interna no coincidía con el producto real, así que esto
    siempre se devuelve para revisión y confirmación del usuario."""
    wb = load_workbook(BytesIO(contenido), data_only=True)
    hoja = wb.active

    filas = list(hoja.iter_rows(values_only=True))
    filas = [f for f in filas if any(c is not None and str(c).strip() for c in f)]

    nombre_base = re.sub(r"\.xlsx?$", "", nombre_archivo or "Receta importada", flags=re.IGNORECASE)
    nombre_base = nombre_base.replace("_", " ").strip() or "Receta importada"

    if not filas:
        return ExcelRecetaImportOut(
            receta_nombre=nombre_base,
            lineas=[],
            mensaje="El archivo está vacío o no se pudo leer. Verifica el formato.",
        )

    # Si la primera fila dice algo como "Receta: Pan de Mantequilla", se usa
    # como sugerencia de nombre -- pero el usuario debe confirmarlo/corregirlo.
    receta_nombre = nombre_base
    primera = filas[0]
    if primera and primera[0] and "receta" in str(primera[0]).lower():
        candidato = primera[-1] if len(primera) > 1 else None
        if candidato and str(candidato).strip():
            receta_nombre = str(candidato).strip()

    # Se busca la fila de encabezado (columnas tipo "insumo"/"ingrediente" + "cantidad")
    idx_header = 0
    for idx, fila in enumerate(filas):
        celdas = [str(c).lower() for c in fila if c is not None]
        if any("insumo" in c or "ingrediente" in c for c in celdas) and any(
            "cantidad" in c for c in celdas
        ):
            idx_header = idx
            break

    header = [str(c).strip().lower() if c is not None else "" for c in filas[idx_header]]
    col_insumo = next((i for i, h in enumerate(header) if "insumo" in h or "ingrediente" in h), 0)
    col_cantidad = next((i for i, h in enumerate(header) if "cantidad" in h), 1)

    insumos_existentes = {i.nombre.strip().lower(): i.nombre for i in db.query(models.Insumo).all()}

    lineas: List[ExcelInsumoLineaOut] = []
    for fila in filas[idx_header + 1 :]:
        if col_insumo >= len(fila) or fila[col_insumo] is None:
            continue
        nombre_insumo = str(fila[col_insumo]).strip()
        if not nombre_insumo:
            continue
        try:
            valor_cantidad = fila[col_cantidad] if col_cantidad < len(fila) else None
            cantidad = float(valor_cantidad) if valor_cantidad is not None else 0
        except (TypeError, ValueError):
            cantidad = 0
        if cantidad <= 0:
            continue
        clave = nombre_insumo.lower()
        existe = clave in insumos_existentes
        lineas.append(
            ExcelInsumoLineaOut(
                insumo_nombre=insumos_existentes.get(clave, nombre_insumo),
                cantidad_requerida=cantidad,
                existe=existe,
            )
        )

    mensaje = "Revisa el nombre de la receta y cada insumo antes de guardar."
    if any(not linea.existe for linea in lineas):
        mensaje += " Hay insumos que no existen todavía en tu inventario; regístralos primero o corrige el nombre."
    if not lineas:
        mensaje = "No se encontraron filas de insumo/cantidad reconocibles en el archivo."

    return ExcelRecetaImportOut(receta_nombre=receta_nombre, lineas=lineas, mensaje=mensaje)
